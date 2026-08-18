"""Suite de aceptacion de ANONIMIZADOR.

Todos los documentos usados son SINTETICOS y se generan al vuelo.
Nunca se usa informacion de una persona real.
"""
from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from anonimizador import casos_sinteticos, config  # noqa: E402
from anonimizador.core import pipeline  # noqa: E402
from anonimizador.core.detectors import clinico  # noqa: E402
from anonimizador.core.validators import alucinaciones, integridad  # noqa: E402
from anonimizador.models import Accion, EstadoRevision, Riesgo  # noqa: E402
from anonimizador.util import workspace  # noqa: E402
from anonimizador.util.hashing import sha256_archivo  # noqa: E402

IDENTIFICADORES_PROHIBIDOS = [
    "Maria Fernanda", "Lopez Quintero", "52.123.456", "52123456",
    "0099887", "320 555 1234", "3205551234", "mafe.lopez1960",
    "Fernando Gil", "Claudia Restrepo", "Sanitas", "San Rafael",
    "hospital-ficticio", "POL-778812", "88123456", "315 444 9988",
]

CLINICOS_OBLIGATORIOS = [
    "8.9 g/dL", "28.4 %", "6 ng/mL", "0.9 mg/dL",
    "colon ascendente", "3.2 x 2.1 cm", "118/76 mmHg",
    "adenocarcinoma", "losartan 50 mg",
]


# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="session")
def muestras(tmp_path_factory):
    destino = tmp_path_factory.mktemp("muestras")
    return {p.name: p for p in casos_sinteticos.generar(destino, forzar=True)}


@pytest.fixture(scope="session")
def resultados(muestras):
    salida = {}
    for nombre, ruta in muestras.items():
        if nombre == "LEEME.txt":
            continue
        salida[nombre] = pipeline.procesar(ruta)
    return salida


def texto_de_salida(resultado):
    """Extrae el texto del archivo generado, con el manejador que corresponda."""
    from anonimizador.formats import registry

    ruta = Path(resultado.archivos["resultado"])
    manejador = registry.manejador_para(ruta.suffix.lower())
    extraido = manejador.reescanear(ruta)
    return "\n".join(u.texto for u in extraido.unidades)


# ---------------------------------------------------------------------------
# TEST 1 - DOCX: identificadores fuera, metadata saneada, clinica intacta
# ---------------------------------------------------------------------------
def test_1_docx_identificadores_metadata_y_clinica(resultados):
    r = resultados["caso_02_historia_clinica.docx"]
    assert r.ok, r.error
    salida = Path(r.archivos["resultado"])
    texto = texto_de_salida(r)

    for identificador in IDENTIFICADORES_PROHIBIDOS:
        assert identificador not in texto, "sobrevivio en el texto: " + identificador

    # tambien en el binario completo del DOCX (XML interno incluido)
    crudo = salida.read_bytes()
    for identificador in IDENTIFICADORES_PROHIBIDOS:
        assert identificador.encode("utf-8") not in crudo, (
            "sobrevivio dentro del ZIP/XML: " + identificador
        )

    for dato in CLINICOS_OBLIGATORIOS[:6]:
        assert dato in texto, "se perdio un dato clinico: " + dato

    # metadata neutra
    import docx

    props = docx.Document(str(salida)).core_properties
    assert props.author == ""
    assert props.last_modified_by == ""
    assert "Lopez" not in (props.title or "")

    # partes peligrosas eliminadas
    nombres = zipfile.ZipFile(salida).namelist()
    for parte in ("word/comments.xml", "docProps/custom.xml", "word/vbaProject.bin"):
        assert parte not in nombres


# ---------------------------------------------------------------------------
# TEST 2 - PDF: no basta con ocultar; la extraccion posterior no recupera nada
# ---------------------------------------------------------------------------
def test_2_pdf_no_recuperable(resultados):
    r = resultados["caso_03_informe_patologia.pdf"]
    assert r.ok, r.error
    salida = Path(r.archivos["resultado"])

    import fitz

    doc = fitz.open(str(salida))
    texto = "".join(doc.load_page(i).get_text() for i in range(doc.page_count))
    meta = " ".join(str(v) for v in (doc.metadata or {}).values())
    doc.close()

    for identificador in IDENTIFICADORES_PROHIBIDOS:
        assert identificador not in texto
        assert identificador not in meta

    crudo = salida.read_bytes()
    for identificador in ["Maria Fernanda", "52.123.456", "mafe.lopez1960"]:
        assert identificador.encode("utf-8") not in crudo

    assert "adenocarcinoma" in texto.lower()
    assert "colon ascendente" in texto.lower()


# ---------------------------------------------------------------------------
# TEST 3 - Hemoglobina: 8.9 g/dL debe salir EXACTAMENTE igual
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("caso", [
    "caso_01_nota_evolucion.txt",
    "caso_02_historia_clinica.docx",
    "caso_03_informe_patologia.pdf",
])
def test_3_valor_de_laboratorio_exacto(resultados, caso):
    texto = texto_de_salida(resultados[caso])
    if "Hemoglobina" in resultados[caso].texto_original:
        assert "8.9 g/dL" in texto
        assert "8.8" not in texto
        assert "9.8" not in texto


def test_3b_integridad_detecta_un_valor_alterado():
    original = "Hemoglobina: 8.9 g/dL. Ferritina: 6 ng/mL."
    manipulado = "Hemoglobina: 8.8 g/dL. Ferritina: 6 ng/mL."
    informe = integridad.comparar(original, manipulado, [])
    assert informe["veredicto"] == "FAIL"
    assert informe["unexpected_change"], "no detecto el cambio de 8.9 a 8.8"
    assert informe["bloquea_aprobacion"] is True


def test_3c_integridad_detecta_un_valor_desaparecido():
    original = "Hemoglobina: 8.9 g/dL. Ferritina: 6 ng/mL."
    manipulado = "Ferritina: 6 ng/mL."
    informe = integridad.comparar(original, manipulado, [])
    assert informe["veredicto"] == "FAIL"
    assert informe["missing"]


# ---------------------------------------------------------------------------
# TEST 4 - Anatomia y lateralidad: colon ascendente NO puede volverse sigmoide
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("caso", [
    "caso_01_nota_evolucion.txt",
    "caso_02_historia_clinica.docx",
    "caso_03_informe_patologia.pdf",
])
def test_4_anatomia_preservada(resultados, caso):
    texto = texto_de_salida(resultados[caso]).lower()
    original = resultados[caso].texto_original.lower()
    if "colon ascendente" in original:
        assert "colon ascendente" in texto
        assert original.count("sigmoide") == texto.count("sigmoide")


def test_4b_cambio_de_anatomia_produce_fail():
    original = "lesion en colon ascendente sin compromiso del recto"
    manipulado = "lesion en colon sigmoide sin compromiso del recto"
    informe = integridad.comparar(original, manipulado, [])
    assert informe["veredicto"] == "FAIL"
    estados = {f["estado"] for f in informe["filas"]}
    assert "MISSING" in estados or "UNEXPECTED_CHANGE" in estados


def test_4c_negacion_preservada():
    original = "No se observan lesiones en colon sigmoide."
    manipulado = "Se observan lesiones en colon sigmoide."
    informe = integridad.comparar(original, manipulado, [])
    assert informe["veredicto"] == "FAIL"


# ---------------------------------------------------------------------------
# TEST 5 - metadata identificadora desaparece
# ---------------------------------------------------------------------------
def test_5_metadatos_desaparecen(resultados):
    for nombre, r in resultados.items():
        adv = r.adversarial
        residual = " ".join(str(v) for v in adv.get("metadatos_residuales", {}).values())
        for identificador in IDENTIFICADORES_PROHIBIDOS:
            assert identificador not in residual, nombre


def test_5b_xlsx_propiedades_neutras(resultados):
    from openpyxl import load_workbook

    r = resultados["caso_05_laboratorios.xlsx"]
    wb = load_workbook(r.archivos["resultado"])
    assert not (wb.properties.creator or "").strip()
    assert "Lopez" not in (wb.properties.title or "")
    assert wb.worksheets[0].title != "HC Lopez Quintero"


# ---------------------------------------------------------------------------
# TEST 6 - matriz de transformacion y auditoria
# ---------------------------------------------------------------------------
def test_6_expediente_de_auditoria_completo(resultados):
    esperados = ["audit.json", "transformation_matrix.csv", "integrity_report.json",
                 "adversarial_scan.json", "audit_report.html", "RESUMEN.md"]
    for nombre, r in resultados.items():
        carpeta = Path(r.archivos["carpeta"])
        for archivo in esperados:
            assert (carpeta / archivo).exists(), nombre + " sin " + archivo
        aud = r.auditoria
        assert aud["archivo_original"]["sha256"]
        assert aud["archivo_resultante"]["sha256"]
        assert aud["riesgo_residual"]["motivos"]
        assert aud["revision_humana"]["estado"] in [e.value for e in EstadoRevision]
        assert len(aud["transformaciones"]["matriz"]) > 0


def test_6b_la_matriz_no_expone_identificadores(resultados):
    import csv

    for nombre, r in resultados.items():
        ruta = Path(r.archivos["matriz_csv"])
        crudo = ruta.read_text(encoding="utf-8-sig")
        for identificador in ["Maria Fernanda Lopez Quintero", "52.123.456",
                              "mafe.lopez1960@correo-ficticio.co"]:
            assert identificador not in crudo, nombre
        filas = list(csv.DictReader(crudo.splitlines(), delimiter=";"))
        assert filas and "accion" in filas[0]


# ---------------------------------------------------------------------------
# TEST 7 - la prueba adversarial se ejecuta de verdad
# ---------------------------------------------------------------------------
def test_7_adversarial_ejecutado(resultados):
    for nombre, r in resultados.items():
        adv = r.adversarial
        assert adv.get("veredicto") in ("PASS", "REVISAR", "FAIL", "NO_EVALUABLE")
        assert adv.get("capas_revisadas"), nombre
        assert adv.get("fugas") == [], nombre + " tiene fugas: " + str(adv.get("fugas"))


def test_7b_adversarial_detecta_una_fuga_inyectada(tmp_path):
    """Si el resultado conserva un identificador, el escaner DEBE cazarlo."""
    from anonimizador.core.validators.adversarial import escanear
    from anonimizador.formats.registry import manejador_para
    from anonimizador.models import Transformacion
    from anonimizador.util.hashing import hash_hallazgo

    filtrado = tmp_path / "con_fuga.txt"
    filtrado.write_text("Paciente: Maria Fernanda Lopez Quintero\nHb 8.9 g/dL\n",
                        encoding="utf-8")
    t = Transformacion(
        categoria="identificador_directo", tipo="nombre_persona", capa="contenido",
        hallazgo_enmascarado="Ma***o", hallazgo_hash=hash_hallazgo("Maria Fernanda Lopez Quintero"),
        accion=Accion.SUSTITUIR, resultado="[PACIENTE]", motivo="prueba",
        integridad_clinica="no aplica", riesgo_residual="bajo",
    )
    informe, alertas = escanear(filtrado, manejador_para(".txt"), [t],
                                ["Maria Fernanda Lopez Quintero"])
    assert informe["veredicto"] == "FAIL"
    assert informe["fugas"]
    assert any(a.codigo == "FUGA_ADVERSARIAL" for a in alertas)


# ---------------------------------------------------------------------------
# El original NUNCA se toca
# ---------------------------------------------------------------------------
def test_original_intacto(muestras, resultados):
    for nombre, r in resultados.items():
        ruta = muestras[nombre]
        assert sha256_archivo(ruta) == r.auditoria["archivo_original"]["sha256"]
        assert r.auditoria["archivo_original"]["intacto_tras_el_proceso"] is True


# ---------------------------------------------------------------------------
# Alucinaciones
# ---------------------------------------------------------------------------
def test_alucinacion_detectada():
    original = "Lesion en colon ascendente. Hemoglobina 8.9 g/dL."
    inventado = ("Lesion en colon ascendente. Hemoglobina 8.9 g/dL. "
                 "Se documenta metastasis hepatica y se inicia oxaliplatino 85 mg.")
    integ = integridad.comparar(original, inventado, [])
    informe = alucinaciones.revisar(original, inventado, integ, [])
    assert informe["veredicto"] == "FAIL"
    assert informe["hay_sospechas"]


def test_sin_alucinacion_en_ejecucion_normal(resultados):
    for nombre, r in resultados.items():
        assert r.integridad["alucinaciones"]["veredicto"] == "PASS", nombre


# ---------------------------------------------------------------------------
# Generalizacion controlada
# ---------------------------------------------------------------------------
def test_edad_se_generaliza_en_banda(resultados):
    texto = texto_de_salida(resultados["caso_01_nota_evolucion.txt"])
    assert "60-69 anos" in texto
    assert "64 anos" not in texto


def test_duracion_clinica_no_se_confunde_con_edad(resultados):
    texto = texto_de_salida(resultados["caso_01_nota_evolucion.txt"])
    # "hace 8 anos" es cronologia clinica: se conserva literal
    assert "hace 8 anos" in texto
    assert "3 meses de evolucion" in texto


def test_fechas_se_vuelven_cronologia_relativa(resultados):
    texto = texto_de_salida(resultados["caso_01_nota_evolucion.txt"])
    assert "Dia 1" in texto
    assert "10/03/2024" not in texto


def test_finalidad_interconsulta_conserva_edad(muestras):
    opciones = config.Opciones(finalidad="interconsulta")
    r = pipeline.procesar(muestras["caso_01_nota_evolucion.txt"], opciones=opciones)
    texto = texto_de_salida(r)
    assert "64 anos" in texto
    assert "Maria Fernanda" not in texto
    # conservar un cuasi-identificador sube el riesgo y lo explica
    assert r.riesgo in (Riesgo.MODERADO, Riesgo.ALTO)


# ---------------------------------------------------------------------------
# Camisa de fuerza clinica
# ---------------------------------------------------------------------------
def test_transformacion_que_pisa_lo_clinico_se_bloquea(tmp_path):
    """Un identificador solapado con un valor clinico NO se transforma a ciegas."""
    from anonimizador.core.detectors import clinico as det

    texto = "Hemoglobina: 8.9 g/dL"
    protegido = det.protegido(texto, 13, 16)
    assert protegido, "el valor clinico deberia estar protegido"


def test_spans_protegidos_cubren_valores_y_terminos():
    texto = "Se observa lesion en colon ascendente. Hemoglobina 8.9 g/dL. TA 118/76 mmHg."
    spans = clinico.spans_protegidos(texto)
    assert spans
    motivos = {m for _, _, m in spans}
    assert "termino clinico" in motivos
    assert "valor con unidad" in motivos


# ---------------------------------------------------------------------------
# Imagenes y QR
# ---------------------------------------------------------------------------
def test_qr_destruido_en_la_imagen(resultados):
    r = resultados["caso_04_etiqueta_estudio.png"]
    assert r.ok
    adv = r.adversarial
    assert adv.get("codigos_graficos_residuales") == [], "el QR sobrevivio"
    salida = Path(r.archivos["resultado"])
    from PIL import Image

    img = Image.open(salida)
    assert not (img.info or {}).get("Author")
    assert "Lopez" not in str(img.info)


def test_imagen_con_texto_no_clasificable_no_se_aprueba_sola(resultados):
    r = resultados["caso_04_etiqueta_estudio.png"]
    assert r.estado != EstadoRevision.APPROVED_BY_HUMAN
    assert r.riesgo in (Riesgo.MODERADO, Riesgo.ALTO, Riesgo.NO_EVALUABLE)


# ---------------------------------------------------------------------------
# Seguridad
# ---------------------------------------------------------------------------
def test_extension_falsificada_se_rechaza(tmp_path):
    falso = tmp_path / "malicioso.docx"
    falso.write_text("esto no es un docx", encoding="utf-8")
    r = pipeline.procesar(falso)
    assert r.ok is False
    assert "no corresponde" in r.error.lower() or "rechaz" in r.error.lower()


def test_extension_no_soportada_se_rechaza(tmp_path):
    falso = tmp_path / "script.exe"
    falso.write_bytes(b"MZ\x90\x00")
    r = pipeline.procesar(falso)
    assert r.ok is False


def test_archivo_vacio_se_rechaza(tmp_path):
    vacio = tmp_path / "vacio.txt"
    vacio.write_bytes(b"")
    r = pipeline.procesar(vacio)
    assert r.ok is False


def test_pdf_escaneado_se_marca_no_evaluable(tmp_path):
    """Un PDF sin texto extraible NO puede pasar por bueno."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    img = tmp_path / "pagina.png"
    im = Image.new("RGB", (800, 1000), (255, 255, 255))
    ImageDraw.Draw(im).text((40, 40), "Paciente: Maria Fernanda Lopez", fill=(0, 0, 0))
    im.save(img)
    pdf = tmp_path / "escaneado.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    c.drawImage(str(img), 20, 20, width=550, height=700)
    c.showPage()
    c.save()

    r = pipeline.procesar(pdf)
    assert r.ok
    assert r.riesgo == Riesgo.NO_EVALUABLE
    assert r.estado == EstadoRevision.REQUIRES_MANUAL_REVIEW
    assert any(a.codigo == "PDF_PROBABLEMENTE_ESCANEADO" for a in r.alertas)


def test_no_hay_red_ni_ia_por_defecto():
    assert config.ALLOW_NETWORK is False
    assert config.ENABLE_LLM is False
    assert config.ENABLE_TELEMETRY is False
    from anonimizador.core.llm import base as capa_llm

    assert capa_llm.obtener_proveedor() is None
    with pytest.raises(capa_llm.LLMDesactivado):
        capa_llm.ProveedorLLM().sugerir_identificadores("hola")


def test_nombre_de_archivo_de_salida_es_neutro(resultados):
    for nombre, r in resultados.items():
        salida = Path(r.archivos["resultado"]).name
        assert "documento_desidentificado" in salida
        for identificador in ["Lopez", "0099887", "52123456"]:
            assert identificador not in salida


def test_nunca_se_afirma_anonimizacion_total(resultados):
    for nombre, r in resultados.items():
        html = Path(r.archivos["html"]).read_text(encoding="utf-8")
        assert "100 % anonimizado" not in html
        assert "100% anonimizado" not in html
        assert "no garantiza" in html
        assert r.estado != EstadoRevision.APPROVED_BY_HUMAN


# ---------------------------------------------------------------------------
# Limpieza del area de trabajo
# ---------------------------------------------------------------------------
def test_eliminar_archivos_de_trabajo():
    workspace.WORKSPACE_DIR.mkdir(parents=True, exist_ok=True)
    prueba = workspace.WORKSPACE_DIR / "_prueba_borrado"
    prueba.mkdir(exist_ok=True)
    (prueba / "x.txt").write_text("temporal", encoding="utf-8")
    workspace.eliminar_archivos_de_trabajo()
    assert not prueba.exists()
